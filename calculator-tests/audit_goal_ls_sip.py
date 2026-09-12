#!/usr/bin/env python3
"""
Audit Unprotected Goal LS–SIP Options v3 workbook structure + engine-parity model.

Excel copy: calculator-tests/Nivra Goal w Current Investment, LS - SIP Options v3.xlsm

Note: the workbook can use a separate current-corpus return (C10) and tax (C11).
The web engine grows current corpus at the goal return/tax (C6/C7). This audit
locks workbook inputs and asserts our engine-equivalent Python model.

  .venv/bin/python calculator-tests/audit_goal_ls_sip.py
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Install openpyxl first: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

XLSX = Path(__file__).resolve().parent / "Nivra Goal w Current Investment, LS - SIP Options v3.xlsm"
SHEET = "Goal Options"


def monthly_rate(annual: float) -> float:
    return (1 + annual) ** (1 / 12) - 1


def fv(rate: float, nper: float, pmt: float, pv: float = 0, typ: int = 1) -> float:
    if rate == 0:
        return -(pv + pmt * nper)
    factor = (1 + rate) ** nper
    if typ == 1:
        return -(pv * factor + pmt * (1 + rate) * (factor - 1) / rate)
    return -(pv * factor + pmt * (factor - 1) / rate)


def pv(rate: float, nper: float, pmt: float, fv_amt: float = 0, typ: int = 0) -> float:
    if rate == 0:
        return -(fv_amt + pmt * nper)
    factor = (1 + rate) ** nper
    if typ == 1:
        return -(fv_amt + pmt * (1 + rate) * (factor - 1) / rate) / factor
    return -(fv_amt + pmt * (factor - 1) / rate) / factor


def inflate(amount: float, inflation: float, years: int) -> float:
    return amount * ((1 + inflation) ** years)


def existing_net_credit(fv_amt: float, invested: float, tax_rate: float) -> float:
    gain = max(0.0, fv_amt - invested)
    return fv_amt - gain * tax_rate


def required_sip(target: float, years: int, annual_return: float, tax_rate: float) -> float:
    """Solve SIP so net-after-tax corpus ≈ target (beginning-of-period)."""
    if target <= 0 or years <= 0:
        return 0.0
    r = monthly_rate(annual_return)
    months = years * 12

    def net_for(monthly: float) -> float:
        maturity = fv(r, months, -monthly, 0, 1)
        invested = monthly * months
        tax = max(0.0, maturity - invested) * tax_rate
        return maturity - tax

    lo, hi = 0.0, max(target, 1.0)
    for _ in range(80):
        if net_for(hi) >= target:
            break
        hi *= 2
    for _ in range(80):
        mid = (lo + hi) / 2
        if net_for(mid) < target:
            lo = mid
        else:
            hi = mid
    return hi


def required_lumpsum(target: float, years: int, annual_return: float, tax_rate: float) -> float:
    if target <= 0 or years <= 0:
        return 0.0

    def net_for(amount: float) -> float:
        maturity = fv(annual_return, years, 0, -amount, 1)
        tax = max(0.0, maturity - amount) * tax_rate
        return maturity - tax

    lo, hi = 0.0, max(target, 1.0)
    for _ in range(80):
        if net_for(hi) >= target:
            break
        hi *= 2
    for _ in range(80):
        mid = (lo + hi) / 2
        if net_for(mid) < target:
            lo = mid
        else:
            hi = mid
    return hi


def model(
    goal_amount: float,
    years: int,
    annual_return: float,
    inflation: float,
    tax_rate: float,
    use_infl_adj: bool,
    current_corpus: float,
    extra_lumpsum: float,
):
    infl_adj = inflate(goal_amount, inflation, years) if inflation else goal_amount
    target = infl_adj if use_infl_adj else goal_amount
    corpus_fv = fv(annual_return, years, 0, -current_corpus, 1)
    credit = existing_net_credit(corpus_fv, current_corpus, tax_rate)
    shortfall = max(0.0, target - credit)
    all_lumpsum = required_lumpsum(shortfall, years, annual_return, tax_rate)
    all_sip = required_sip(shortfall, years, annual_return, tax_rate)

    mix_corpus = current_corpus + extra_lumpsum
    mix_fv = fv(annual_return, years, 0, -mix_corpus, 1)
    mix_credit = existing_net_credit(mix_fv, mix_corpus, tax_rate)
    mix_shortfall = max(0.0, target - mix_credit)
    mix_sip = required_sip(mix_shortfall, years, annual_return, tax_rate)
    extra_fv = fv(annual_return, years, 0, -extra_lumpsum, 1) if extra_lumpsum else 0.0

    return {
        "inflAdjGoal": infl_adj,
        "targetGoal": target,
        "existingCredit": credit,
        "shortfall": shortfall,
        "allLumpsum": all_lumpsum,
        "allSip": all_sip,
        "mixSip": mix_sip,
        "mixShortfall": mix_shortfall,
        "extraLumpsumFv": extra_fv,
    }


def close(a: float, b: float, rel: float = 1e-8) -> bool:
    return abs(a - b) <= max(1e-4, abs(b) * rel)


def main() -> int:
    if not XLSX.exists():
        print(f"Missing workbook: {XLSX}", file=sys.stderr)
        return 1

    wb = load_workbook(XLSX, data_only=False)
    assert SHEET in wb.sheetnames, f"expected sheet {SHEET}"
    ws = wb[SHEET]

    assert ws["C4"].value == 10_000_000
    assert ws["C5"].value == 10
    assert abs(float(ws["C6"].value) - 0.12) < 1e-12
    assert abs(float(ws["C7"].value) - 0.125) < 1e-12
    assert ws["C9"].value == 500_000
    assert ws["C13"].value == 100_000
    print("OK workbook inputs C4/C5/C6/C7/C9/C13")

    # Engine-aligned sample used by TS parity (15y infl-adj golden path)
    sample = model(
        goal_amount=10_000_000,
        years=15,
        annual_return=0.12,
        inflation=0.0525,
        tax_rate=0.125,
        use_infl_adj=True,
        current_corpus=500_000,
        extra_lumpsum=200_000,
    )
    # Locked against packages/finance calculateGoalLsSipOptions golden
    expect = {
        "allLumpsum": 3883931.0297774756,
        "allSip": 43484.474879902504,
        "mixSip": 41245.27575167319,
    }
    for key, exp in expect.items():
        got = sample[key]
        assert close(got, exp), f"{key}: expected {exp}, got {got}"
        print(f"OK engine-model {key}={got}")

    # Workbook default path (10y, no infl adj, extra 1L) — engine rates only
    book = model(
        goal_amount=10_000_000,
        years=10,
        annual_return=0.12,
        inflation=0.0,
        tax_rate=0.125,
        use_infl_adj=False,
        current_corpus=500_000,
        extra_lumpsum=100_000,
    )
    assert book["targetGoal"] == 10_000_000
    assert book["allSip"] > book["mixSip"] > 0
    assert book["allLumpsum"] > 0
    print("OK workbook-default engine path allSip > mixSip")
    print("NOTE Excel O5 uses current-corpus rate C10/C11; web engine uses C6/C7 for corpus too.")
    print("All audits passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
