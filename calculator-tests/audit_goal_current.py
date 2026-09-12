#!/usr/bin/env python3
"""
Audit Unprotected Goal with Current Investment workbook + engine-parity model.

Excel copy: calculator-tests/Nivra Goal with Current Investment - LS, SIP, SU_SIP.xlsm

Note: the workbook can use a separate current-SIP return (D11). The web engine
grows current corpus + current SIP at the goal return/tax (C6/C7). This audit
locks workbook inputs and asserts our engine-equivalent Python model.

  .venv/bin/python calculator-tests/audit_goal_current.py
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Install openpyxl first: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

XLSX = Path(__file__).resolve().parent / "Nivra Goal with Current Investment - LS, SIP, SU_SIP.xlsm"
SHEET = "Goal Options"


def monthly_rate(annual: float) -> float:
    return (1 + annual) ** (1 / 12) - 1


def fv(rate: float, nper: float, pmt: float, pv: float = 0, typ: int = 1) -> float:
    if rate == 0:
        return -(pv + pmt * nper)
    factor = (1 + rate) ** nper
    if typ == 1:
        return -(pv * factor + pmt * (1 + rate) * (factor - 1) / rate)
    return -(pv * factor + pmt * (factor - 1) / rate)


def inflate(amount: float, inflation: float, years: int) -> float:
    return amount * ((1 + inflation) ** years)


def existing_net_credit(fv_amt: float, invested: float, tax_rate: float) -> float:
    gain = max(0.0, fv_amt - invested)
    return fv_amt - gain * tax_rate


def required_sip(target: float, years: int, annual_return: float, tax_rate: float) -> float:
    if target <= 0 or years <= 0:
        return 0.0
    r = monthly_rate(annual_return)
    months = years * 12

    def net_for(monthly: float) -> float:
        maturity = fv(r, months, -monthly, 0, 1)
        invested = monthly * months
        tax = max(0.0, maturity - invested) * tax_rate
        return maturity - tax

    lo, hi = 0.0, max(target, 1.0)
    for _ in range(80):
        if net_for(hi) >= target:
            break
        hi *= 2
    for _ in range(80):
        mid = (lo + hi) / 2
        if net_for(mid) < target:
            lo = mid
        else:
            hi = mid
    return hi


def required_lumpsum(target: float, years: int, annual_return: float, tax_rate: float) -> float:
    if target <= 0 or years <= 0:
        return 0.0

    def net_for(amount: float) -> float:
        maturity = fv(annual_return, years, 0, -amount, 1)
        tax = max(0.0, maturity - amount) * tax_rate
        return maturity - tax

    lo, hi = 0.0, max(target, 1.0)
    for _ in range(80):
        if net_for(hi) >= target:
            break
        hi *= 2
    for _ in range(80):
        mid = (lo + hi) / 2
        if net_for(mid) < target:
            lo = mid
        else:
            hi = mid
    return hi


def model(
    goal_amount: float,
    years: int,
    annual_return: float,
    inflation: float,
    tax_rate: float,
    use_infl_adj: bool,
    current_corpus: float,
    current_monthly_sip: float,
):
    infl_adj = inflate(goal_amount, inflation, years) if inflation else goal_amount
    target = infl_adj if use_infl_adj else goal_amount
    corpus_fv = fv(annual_return, years, 0, -current_corpus, 1) if current_corpus else 0.0
    r = monthly_rate(annual_return)
    sip_fv = (
        fv(r, years * 12, -current_monthly_sip, 0, 1) if current_monthly_sip else 0.0
    )
    existing_fv = corpus_fv + sip_fv
    existing_invested = current_corpus + current_monthly_sip * years * 12
    credit = existing_net_credit(existing_fv, existing_invested, tax_rate)
    shortfall = max(0.0, target - credit)
    return {
        "inflAdjGoal": infl_adj,
        "targetGoal": target,
        "netCredit": credit,
        "shortfall": shortfall,
        "allLumpsum": required_lumpsum(shortfall, years, annual_return, tax_rate),
        "allSip": required_sip(shortfall, years, annual_return, tax_rate),
    }


def close(a: float, b: float, rel: float = 1e-8) -> bool:
    return abs(a - b) <= max(1e-4, abs(b) * rel)


def main() -> int:
    if not XLSX.exists():
        print(f"Missing workbook: {XLSX}", file=sys.stderr)
        return 1

    wb = load_workbook(XLSX, data_only=False)
    assert SHEET in wb.sheetnames, f"expected sheet {SHEET}"
    ws = wb[SHEET]
    assert ws["B1"].value and "Current Investment" in str(ws["B1"].value)
    assert ws["C4"].value == 10_000_000
    assert ws["C5"].value == 15
    assert abs(float(ws["C6"].value) - 0.12) < 1e-12
    assert abs(float(ws["C7"].value) - 0.125) < 1e-12
    assert abs(float(ws["C25"].value) - 0.1) < 1e-12
    print("OK workbook title + C4/C5/C6/C7/C25")

    sample = model(
        goal_amount=10_000_000,
        years=15,
        annual_return=0.12,
        inflation=0.0525,
        tax_rate=0.125,
        use_infl_adj=True,
        current_corpus=500_000,
        current_monthly_sip=5_000,
    )
    expect = {
        "shortfall": 16892374.415978163,
        "allSip": 38484.4748799025,
        "allLumpsum": 3437342.7887438303,
    }
    for key, exp in expect.items():
        got = sample[key]
        assert close(got, exp), f"{key}: expected {exp}, got {got}"
        print(f"OK engine-model {key}={got}")

    print("All audits passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
