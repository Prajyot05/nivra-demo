#!/usr/bin/env python3
"""
Audit Unprotected Nivra SIP Step-Up Calculator v1.xlsm and assert formula parity.

  .venv/bin/python calculator-tests/audit_sip_stepup.py

Note: Excel C8 (Total Investment Duration) can exceed SIP years; the web
growth-stepup API currently uses sipYears as the horizon (equal-years path).
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Install openpyxl first: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

XLSX = Path(__file__).resolve().parent / "Nivra SIP Step-Up Calculator v1.xlsm"
SHEET = "SIP Calculator"


def monthly_rate(annual: float) -> float:
    return (1 + annual) ** (1 / 12) - 1


def fv(rate: float, nper: float, pmt: float, pv: float = 0, typ: int = 1) -> float:
    if rate == 0:
        return -(pv + pmt * nper)
    factor = (1 + rate) ** nper
    if typ == 1:
        return -(pv * factor + pmt * (1 + rate) * (factor - 1) / rate)
    return -(pv * factor + pmt * (factor - 1) / rate)


def year_index_for_month(month_number: int) -> int:
    if month_number % 12 == 0:
        return month_number // 12 - 1
    return month_number // 12


def step_up_monthly(start: float, step_up: float, month_number: int) -> float:
    return start * (1 + step_up) ** year_index_for_month(month_number)


def model(
    start_monthly: float,
    sip_years: int,
    annual_return: float,
    step_up: float,
    inflation: float,
    tax_rate: float = 0.0,
    invest_years: int | None = None,
):
    horizon = invest_years if invest_years is not None else sip_years
    r = monthly_rate(annual_return)
    n = sip_years * 12
    maturity = 0.0
    invested = 0.0
    end_monthly = 0.0
    for month in range(1, n + 1):
        monthly = step_up_monthly(start_monthly, step_up, month)
        end_monthly = monthly
        invested += monthly
        months_remaining = n - month + 1
        maturity += fv(r, months_remaining, 0, -monthly, 1)
    extra = horizon - sip_years
    if extra > 0:
        maturity = fv(annual_return, extra, 0, -maturity, 1)
    gain = maturity - invested
    tax = max(0.0, gain) * tax_rate
    infl = maturity if inflation == 0 else maturity / ((1 + inflation) ** sip_years)
    return {
        "maturity": maturity,
        "totalInvested": invested,
        "gain": gain,
        "endMonthly": end_monthly,
        "inflationAdjusted": infl,
        "tax": tax,
        "netAfterTax": maturity - tax,
    }


def close(a: float, b: float, rel: float = 1e-9) -> bool:
    return abs(a - b) <= max(1e-6, abs(b) * rel)


def main() -> int:
    if not XLSX.exists():
        print(f"Missing workbook: {XLSX}", file=sys.stderr)
        return 1

    wb = load_workbook(XLSX, data_only=False, keep_vba=True)
    assert SHEET in wb.sheetnames, wb.sheetnames
    ws = wb[SHEET]

    print("=== SIP Step-Up Calculator v1 — field audit ===\n")
    editable = {
        "C4": "Monthly Investment Amount (start)",
        "C5": "Step-Up % per year",
        "C6": "SIP Term in Years",
        "C7": "Expected Rate of Return",
        "C8": "Total Investment Duration in Years",
        "C13": "Inflation Rate/Year",
        "C16": "Delay in Investing - Months",
        "H4": "Name",
        "H5": "Age",
    }
    print("Editable inputs:")
    for addr, label in editable.items():
        print(f"  {addr}: {label} = {ws[addr].value!r} locked={ws[addr].protection.locked}")

    print("\nComputed (formulas):")
    for addr in ["C10", "C11", "C14", "C17", "C18", "C20", "C21", "D8", "W1", "W4"]:
        print(f"  {addr}: {ws[addr].value!r}")

    print("\nValidation: D8 = Invalid when C8 < SIPYears")

    cached = load_workbook(XLSX, data_only=True, keep_vba=True)[SHEET]
    sample_equal = model(5000, 10, 0.12, 0.1, 0.0575)
    sample_15 = model(5000, 10, 0.12, 0.1, 0.0575, invest_years=15)

    print("\n=== Equal-years web/API sample (₹5k start, 10% step-up, 10y, 12%) ===")
    for k, v in sample_equal.items():
        print(f"  {k}: {v}")

    print("\n=== Cached Excel vs model (workbook C8 horizon) ===")
    horizon = int(cached["C8"].value or 10)
    expected = sample_15 if horizon == 15 else sample_equal
    checks = [
        ("C20", expected["totalInvested"]),
        ("C10", expected["maturity"]),
        ("C21", expected["gain"]),
        ("C14", expected["inflationAdjusted"]),
    ]
    failed = 0
    for addr, exp in checks:
        actual = cached[addr].value
        ok = actual is not None and close(float(actual), float(exp))
        status = "OK" if ok else "FAIL"
        if not ok:
            failed += 1
        print(f"  {status} {addr}: excel={actual!r} model={exp} (horizon={horizon})")

    print("\n=== Golden equal-years parity ===")
    assert close(sample_equal["maturity"], 1634449.2409538408)
    assert close(sample_equal["totalInvested"], 956245.4760600011)
    assert close(sample_equal["endMonthly"], 5000 * 1.1**9)
    print("  OK golden sample numbers")

    if failed:
        print(f"\n{failed} check(s) failed", file=sys.stderr)
        return 1
    print("\nAll checks passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
