#!/usr/bin/env python3
"""
Audit Unprotected Nivra SIP Calculator v3.xlsm and assert formula parity.

  .venv/bin/python calculator-tests/audit_sip_calculator.py
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Install openpyxl first: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

XLSX = Path(__file__).resolve().parent / "Nivra SIP Calculator v3.xlsm"
SHEET = "SIP Calculator"


def monthly_rate(annual: float) -> float:
    return (1 + annual) ** (1 / 12) - 1


def fv(rate: float, nper: float, pmt: float, pv: float = 0, typ: int = 1) -> float:
    """Excel FV."""
    if rate == 0:
        return -(pv + pmt * nper)
    factor = (1 + rate) ** nper
    if typ == 1:
        return -(pv * factor + pmt * (1 + rate) * (factor - 1) / rate)
    return -(pv * factor + pmt * (factor - 1) / rate)


def model(
    monthly: float,
    sip_years: int,
    invest_years: int,
    annual_return: float,
    inflation: float,
    delay_months: float,
    tax_rate: float = 0.0,
):
    sip_months = sip_years * 12
    r = monthly_rate(annual_return)
    sip_end = fv(r, sip_months, -monthly, 0, 1)
    extra = invest_years - sip_years
    maturity = fv(annual_return, extra, 0, -sip_end, 1) if extra else sip_end
    invested = monthly * sip_months
    delayed = None
    cost = None
    if delay_months > 0:
        invest_months = invest_years * 12
        if invest_months <= delay_months:
            delayed = 0.0
        elif invest_months <= sip_months + delay_months:
            delayed = fv(r, invest_months - delay_months, -monthly, 0, 1)
        else:
            delayed = fv(r, invest_months - sip_months - delay_months, 0, -sip_end, 1)
        cost = maturity - delayed
    infl_adj = maturity if inflation == 0 else maturity / ((1 + inflation) ** sip_years)
    gain = maturity - invested
    tax = max(0.0, gain) * tax_rate
    return {
        "maturity": maturity,
        "totalInvested": invested,
        "gain": gain,
        "inflationAdjusted": infl_adj,
        "delayedMaturity": delayed,
        "costOfDelay": cost,
        "tax": tax,
        "netAfterTax": maturity - tax,
    }


def close(a: float, b: float, rel: float = 1e-9) -> bool:
    return abs(a - b) <= max(1e-6, abs(b) * rel)


def main() -> int:
    if not XLSX.exists():
        print(f"Missing workbook: {XLSX}", file=sys.stderr)
        return 1

    wb = load_workbook(XLSX, data_only=False, keep_vba=True)
    assert SHEET in wb.sheetnames, wb.sheetnames
    ws = wb[SHEET]

    print("=== SIP Calculator v3 — field audit ===\n")
    editable = {
        "C4": "Monthly Investment Amount",
        "C5": "SIP Term in Years",
        "C6": "Expected Rate of Return",
        "C7": "Total Investment Duration in Years",
        "C12": "Inflation Rate/Year",
        "C15": "Delay in Investing - Months",
        "H4": "Name",
        "H5": "Age",
    }
    print("Editable inputs:")
    for addr, label in editable.items():
        print(f"  {addr}: {label} = {ws[addr].value!r} locked={ws[addr].protection.locked}")

    print("\nComputed (formulas):")
    for addr in ["C9", "C10", "C13", "C16", "C17", "C19", "C20", "D7", "X1", "X4", "X5", "X6"]:
        print(f"  {addr}: {ws[addr].value!r}")

    print("\nValidation: D7 = Invalid when C7 < SIPYears")

    cached = load_workbook(XLSX, data_only=True, keep_vba=True)[SHEET]
    # Workbook default is often 40y horizon; engine sample uses equal SIP/horizon.
    sample_5 = model(1500, 5, 5, 0.12, 0.0575, 6)
    sample_40 = model(1500, 5, 40, 0.12, 0.0575, 6)

    print("\n=== Model sample (₹1500 / 5y SIP / 5y horizon / 12% / infl 5.75% / delay 6m) ===")
    for k, v in sample_5.items():
        print(f"  {k}: {v}")

    print("\n=== Cached Excel vs model (uses workbook C7 horizon) ===")
    horizon = int(cached["C7"].value or 40)
    expected = sample_40 if horizon == 40 else sample_5
    checks = [
        ("C19", expected["totalInvested"]),
        ("C9", expected["maturity"]),
        ("C20", expected["gain"]),
        ("C13", expected["inflationAdjusted"]),
        ("C16", expected["delayedMaturity"]),
        ("C17", expected["costOfDelay"]),
    ]
    failed = 0
    for addr, exp in checks:
        actual = cached[addr].value
        ok = actual is not None and exp is not None and close(float(actual), float(exp))
        status = "OK" if ok else "FAIL"
        if not ok:
            failed += 1
        print(f"  {status} {addr}: excel={actual!r} model={exp} (horizon={horizon})")

    print("\n=== Equal-horizon 5y parity (golden / UI default) ===")
    assert close(sample_5["maturity"], 121655.41880029078)
    assert close(sample_5["totalInvested"], 90_000)
    assert close(sample_5["delayedMaturity"] or 0, 106162.42472526057)
    print("  OK golden sample numbers")

    if failed:
        print(f"\n{failed} check(s) failed", file=sys.stderr)
        return 1
    print("\nAll checks passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
