#!/usr/bin/env python3
"""
Audit Unprotected Nivra One-Time Investment v2.xlsm and assert formula parity.

  .venv/bin/python calculator-tests/audit_one_time_investment.py
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Install openpyxl first: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

XLSX = Path(__file__).resolve().parent / "Nivra One-Time Investment v2.xlsm"
SHEET = "One-Time Investment"


def fv(rate: float, nper: float, pmt: float, pv: float, typ: int = 1) -> float:
    """Excel FV with type=1 (annuity due). For lumpsum pmt=0 → (1+r)^n * -pv."""
    if pmt == 0:
        return -pv * (1 + rate) ** nper
    # unused for this sheet
    raise NotImplementedError


def model(amount: float, years: float, rate: float, inflation: float, delay_months: float):
    maturity = fv(rate, years, 0, -amount, 1)
    inflation_adj = maturity / ((1 + inflation) ** years)
    delayed = fv(rate, years - delay_months / 12, 0, -amount, 1) if delay_months > 0 else None
    cost = (maturity - delayed) if delayed is not None else None
    return {
        "maturity": maturity,
        "inflationAdjusted": inflation_adj,
        "gain": maturity - amount,
        "inflationAdjustedGain": inflation_adj - amount,
        "delayedMaturity": delayed,
        "costOfDelay": cost,
        "schedule": [
            {
                "year": y,
                "yearEnd": amount * (1 + rate) ** y,
                "inflationAdjusted": (amount * (1 + rate) ** y) / ((1 + inflation) ** y),
            }
            for y in range(1, int(years) + 1)
        ],
    }


def close(a: float, b: float, rel: float = 1e-9) -> bool:
    return abs(a - b) <= max(1e-6, abs(b) * rel)


def main() -> int:
    if not XLSX.exists():
        print(f"Missing workbook: {XLSX}", file=sys.stderr)
        return 1

    wb = load_workbook(XLSX, data_only=False, keep_vba=True)
    assert SHEET in wb.sheetnames, wb.sheetnames
    ws = wb[SHEET]

    print("=== One-Time Investment v2 — field audit ===\n")
    editable = {
        "C4": "Investment Amount",
        "C5": "Term in Years",
        "C6": "Expected Rate of Return",
        "C11": "Inflation Rate/Year",
        "C17": "Delay in Investing - Months",
        "H4": "Name",
        "H5": "Age",
    }
    print("Editable inputs:")
    for addr, label in editable.items():
        print(f"  {addr}: {label} = {ws[addr].value!r} locked={ws[addr].protection.locked}")

    print("\nComputed (formulas):")
    for addr in ["C8", "C12", "C14", "C15", "C18", "C19"]:
        print(f"  {addr}: {ws[addr].value!r}")

    cached = load_workbook(XLSX, data_only=True, keep_vba=True)[SHEET]
    sample = model(5_000_000, 16, 0.12, 0.0575, 6)
    checks = [
        ("C8", sample["maturity"]),
        ("C12", sample["inflationAdjusted"]),
        ("C14", sample["gain"]),
        ("C15", sample["inflationAdjustedGain"]),
        ("C18", sample["delayedMaturity"]),
        ("C19", sample["costOfDelay"]),
    ]
    print("\n=== Sample parity (₹50L / 16y / 12% / infl 5.75% / delay 6m) ===")
    failed = 0
    for addr, expected in checks:
        actual = cached[addr].value
        ok = actual is not None and close(float(actual), float(expected))
        if not ok:
            failed += 1
        print(f"  {'OK' if ok else 'FAIL'} {addr}: excel={actual!r} model={expected!r}")

    print("\n=== Schedule year 16 ===")
    y16 = sample["schedule"][-1]
    print(f"  yearEnd={y16['yearEnd']!r} inflationAdj={y16['inflationAdjusted']!r}")
    assert close(y16["yearEnd"], sample["maturity"])

    if failed:
        print(f"\nFAILED {failed} checks", file=sys.stderr)
        return 1
    print("\nAll cached sample checks passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
