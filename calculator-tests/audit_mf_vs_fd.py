#!/usr/bin/env python3
"""
Audit Unprotected Nivra MF vs FD v1.xlsm and assert formula parity.

Requires: openpyxl (use a venv if needed).
  python3 -m venv .venv && .venv/bin/pip install openpyxl
  .venv/bin/python calculator-tests/audit_mf_vs_fd.py
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Install openpyxl first: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

XLSX = Path(__file__).resolve().parent / "Nivra MF vs FD v1.xlsm"
SHEET = "MF vs FD"

# Excel ActiveX ComboBox option ranges (from VML FmlaRange)
MF_TAX_OPTIONS = (0.10, 0.125, 0.20, 0.30)  # J2:J5
FD_TAX_OPTIONS = (0.20, 0.25, 0.30)  # K2:K4

# Unlocked input cells (sheet protection off in this file; locked=False = editable intent)
EDITABLE = {
    "C3": "Investment Period in Days (shared; F3 mirrors)",
    "C4": "MF Interest Amount (annual rate)",
    "F4": "FD Interest Amount (annual rate)",
    "C5": "Investment Amount (shared; F5 mirrors)",
    "C13": "MF Tax Rate (also set by cbMFTax → J10)",
    "F13": "FD Tax Rate (also set by cbFDTax → K10)",
}

# Formula / mirror outputs — not user inputs
COMPUTED = {
    "F3": "=C3",
    "F5": "=C5",
    "C8": "=C5*C4",
    "F8": "=F5*F4",
    "C9": "=C8/365",
    "F9": "=F8/365",
    "C10": "=C9*C3",
    "F10": "=F9*F3",
    "C14": "=C10-(C10*C13)",
    "F14": "=F10-(F10*F13)",
    "C15": "=IF(C14-F14<=0,0,C14-F14)",
    "F15": "=IF(F14-C14<=0,0,F14-C14)",
}

VBA = """
Private Sub cbFDTax_Change()
    Range("F13") = Range("K10").Value2
End Sub
Private Sub cbMFTax_Change()
    Range("C13") = Range("J10").Value2
End Sub
"""


def excel_leg(amount: float, days: float, rate: float, tax: float) -> dict[str, float]:
    """Mirror sheet formulas exactly (365-day count, simple interest)."""
    annualized = amount * rate
    per_day = annualized / 365
    pre_tax = per_day * days
    tax_amt = pre_tax * tax
    post_tax = pre_tax - tax_amt
    return {
        "annualizedReturn": annualized,
        "returnPerDay": per_day,
        "preTax": pre_tax,
        "tax": tax_amt,
        "postTax": post_tax,
        "invested": amount,
        "gain": pre_tax,
        "net": amount + post_tax,
    }


def compare(amount: float, days: float, mf_rate: float, fd_rate: float, mf_tax: float, fd_tax: float):
    mf = excel_leg(amount, days, mf_rate, mf_tax)
    fd = excel_leg(amount, days, fd_rate, fd_tax)
    return {
        "mf": mf,
        "fd": fd,
        "mfAdvantage": max(0.0, mf["postTax"] - fd["postTax"]),
        "fdAdvantage": max(0.0, fd["postTax"] - mf["postTax"]),
        "difference": mf["postTax"] - fd["postTax"],
    }


def close(a: float, b: float, rel: float = 1e-12) -> bool:
    return abs(a - b) <= max(1e-9, abs(b) * rel)


def main() -> int:
    if not XLSX.exists():
        print(f"Missing workbook: {XLSX}", file=sys.stderr)
        return 1

    wb = load_workbook(XLSX, data_only=False, keep_vba=True)
    assert wb.sheetnames == [SHEET], wb.sheetnames
    ws = wb[SHEET]

    print("=== Nivra MF vs FD v1 — field audit ===\n")
    print("Sheet:", SHEET)
    print("Selectable (ActiveX ComboBox):")
    print(f"  cbMFTax → options J2:J5 = {[f'{x*100:g}%' for x in MF_TAX_OPTIONS]} → writes C13")
    print(f"  cbFDTax → options K2:K4 = {[f'{x*100:g}%' for x in FD_TAX_OPTIONS]} → writes F13")
    print("\nEditable inputs:")
    for addr, desc in EDITABLE.items():
        print(f"  {addr}: {desc}  (sample={ws[addr].value!r}, locked={ws[addr].protection.locked})")
    print("\nLocked mirrors / outputs:")
    for addr, formula in COMPUTED.items():
        print(f"  {addr}: {formula}  (cell={ws[addr].value!r})")

    # Combo option cells
    mf_opts = [ws[f"J{r}"].value for r in range(2, 6)]
    fd_opts = [ws[f"K{r}"].value for r in range(2, 5)]
    assert mf_opts == list(MF_TAX_OPTIONS), mf_opts
    assert fd_opts == list(FD_TAX_OPTIONS), fd_opts

    # Data validation leftover on D13/G13 (display helpers; formulas use C13/F13)
    dvs = list(ws.data_validations.dataValidation) if ws.data_validations else []
    print("\nData validations (legacy list on D13/G13):")
    for dv in dvs:
        print(f"  {dv.sqref}: {dv.formula1}")

    print("\nVBA (extracted):")
    print(VBA)

    # Cached sample
    cached = load_workbook(XLSX, data_only=True, keep_vba=True)[SHEET]
    sample = compare(1_000_000_000, 15, 0.05, 0.03, 0.20, 0.25)
    checks = [
        ("C8", sample["mf"]["annualizedReturn"]),
        ("F8", sample["fd"]["annualizedReturn"]),
        ("C9", sample["mf"]["returnPerDay"]),
        ("F9", sample["fd"]["returnPerDay"]),
        ("C10", sample["mf"]["preTax"]),
        ("F10", sample["fd"]["preTax"]),
        ("C14", sample["mf"]["postTax"]),
        ("F14", sample["fd"]["postTax"]),
        ("C15", sample["mfAdvantage"]),
        ("F15", sample["fdAdvantage"]),
    ]
    print("=== Sample parity (₹100 Cr / 15d / 5% vs 3% / 20% vs 25% tax) ===")
    failed = 0
    for addr, expected in checks:
        actual = cached[addr].value
        ok = actual is not None and close(float(actual), expected)
        status = "OK" if ok else "FAIL"
        if not ok:
            failed += 1
        print(f"  {status} {addr}: excel={actual!r} model={expected!r}")

    # Extra scenarios (formula model only — same as engine)
    scenarios = [
        ("short 7d", 5_000_000, 7, 0.06, 0.055, 0.125, 0.20),
        ("hnw 90d", 5_000_000_000, 90, 0.08, 0.065, 0.20, 0.30),
        ("fd wins", 1_000_000, 30, 0.04, 0.07, 0.30, 0.20),
        ("mf tax 10%", 10_000_000, 45, 0.05, 0.03, 0.10, 0.25),
        ("mf tax 12.5%", 10_000_000, 45, 0.05, 0.03, 0.125, 0.25),
    ]
    print("\n=== Scenario smoke (model self-consistency) ===")
    for name, *args in scenarios:
        r = compare(*args)
        assert r["mfAdvantage"] >= 0 and r["fdAdvantage"] >= 0
        assert abs(r["difference"] - (r["mf"]["postTax"] - r["fd"]["postTax"])) < 1e-9
        assert (r["mfAdvantage"] > 0) + (r["fdAdvantage"] > 0) <= 1 or abs(r["difference"]) < 1e-9
        print(
            f"  {name}: MF post={r['mf']['postTax']:.4f} FD post={r['fd']['postTax']:.4f} "
            f"mfAdv={r['mfAdvantage']:.4f} fdAdv={r['fdAdvantage']:.4f}"
        )

    # Web field checklist
    print("\n=== Required web inputs (must all exist) ===")
    required = [
        "Investment Period in Days",
        "MF Interest / return %",
        "FD Interest / return %",
        "Investment Amount",
        "MF Tax Rate select [10, 12.5, 20, 30]%",
        "FD Tax Rate select [20, 25, 30]%",
    ]
    for item in required:
        print(f"  [ ] {item}")
    print("\n=== Required web outputs ===")
    for item in [
        "Annualized Return (MF + FD)",
        "Return Per Day (MF + FD)",
        "Expected Pre-Tax Return (MF + FD)",
        "Tax amount (derived)",
        "Post-Tax Return (MF + FD)",
        "Difference in Return (MF adv + FD adv)",
        "Excel educational notes (FD penalty + tenure flexibility)",
    ]:
        print(f"  [ ] {item}")

    if failed:
        print(f"\nFAILED {failed} cached-value checks", file=sys.stderr)
        return 1
    print("\nAll cached sample checks passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
