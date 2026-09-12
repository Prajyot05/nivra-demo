#!/usr/bin/env python3
"""
Audit Full Set Multiple Goals with Corpus Assignment v2 workbook + engine-parity model.

Excel copy: calculator-tests/Nivra Multiple Goals with Corpus Assignment v2.xlsm

Sheets:
  Goal Calculator  — zero-corpus Goal plan sample (Education / House1 / House2 / Car / Marriage)
  GCwithCorpus     — corpus assignment path (inputs locked; engine assigns soonest-first)

Note: Excel GCwithCorpus can use a separate corpus return (M7). The web engine assigns
corpus by PV of required lumpsum at the goal ST/LT yield, then solves remaining SIP/LS.

  .venv/bin/python calculator-tests/audit_multi_goal_assign.py
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Install openpyxl first: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

XLSX = Path(__file__).resolve().parent / "Nivra Multiple Goals with Corpus Assignment v2.xlsm"


def monthly_rate(annual: float) -> float:
    return (1 + annual) ** (1 / 12) - 1


def fv(rate: float, nper: float, pmt: float, pv: float = 0, typ: int = 1) -> float:
    if rate == 0:
        return -(pv + pmt * nper)
    factor = (1 + rate) ** nper
    if typ == 1:
        return -(pv * factor + pmt * (1 + rate) * (factor - 1) / rate)
    return -(pv * factor + pmt * (factor - 1) / rate)


def inflate(amount: float, inflation: float, years: int) -> float:
    return amount * ((1 + inflation) ** years)


def required_sip(target: float, years: float, annual_return: float, tax_rate: float) -> float:
    if target <= 0 or years <= 0:
        return 0.0
    r = monthly_rate(annual_return)
    months = years * 12

    def net_for(monthly: float) -> float:
        maturity = fv(r, months, -monthly, 0, 1)
        invested = monthly * months
        tax = max(0.0, maturity - invested) * tax_rate
        return maturity - tax

    lo, hi = 0.0, max(target, 1.0)
    for _ in range(80):
        if net_for(hi) >= target:
            break
        hi *= 2
    for _ in range(80):
        mid = (lo + hi) / 2
        if net_for(mid) < target:
            lo = mid
        else:
            hi = mid
    return hi


def required_lumpsum(target: float, years: float, annual_return: float, tax_rate: float) -> float:
    if target <= 0 or years <= 0:
        return 0.0

    def net_for(amount: float) -> float:
        maturity = fv(annual_return, years, 0, -amount, 1)
        tax = max(0.0, maturity - amount) * tax_rate
        return maturity - tax

    lo, hi = 0.0, max(target, 1.0)
    for _ in range(80):
        if net_for(hi) >= target:
            break
        hi *= 2
    for _ in range(80):
        mid = (lo + hi) / 2
        if net_for(mid) < target:
            lo = mid
        else:
            hi = mid
    return hi


def yield_for(years: int, st_years: int, st: float, lt: float) -> float:
    return st if years <= st_years else lt


def model(
    short_term_years: int,
    short_term_return: float,
    long_term_return: float,
    inflation: float,
    tax_rate: float,
    current_corpus: float,
    goals: list[dict],
):
    remaining = current_corpus
    indexed = [
        {**g, "index": i}
        for i, g in enumerate(goals)
        if g["amount"] > 0 and g["years"] > 0
    ]
    ordered = sorted(indexed, key=lambda g: (g["years"], g["index"]))
    assigned_by_index: dict[int, float] = {}
    for g in ordered:
        infl = inflate(g["amount"], inflation, g["years"])
        yld = yield_for(g["years"], short_term_years, short_term_return, long_term_return)
        full_ls = required_lumpsum(infl, g["years"], yld, tax_rate)
        assigned = min(remaining, full_ls)
        assigned_by_index[g["index"]] = assigned
        remaining -= assigned

    out_goals = []
    for i, g in enumerate(goals):
        if g["amount"] <= 0 or g["years"] <= 0:
            out_goals.append({**g, "monthlySip": 0.0, "lumpsum": 0.0, "assigned": 0.0})
            continue
        infl = inflate(g["amount"], inflation, g["years"])
        yld = yield_for(g["years"], short_term_years, short_term_return, long_term_return)
        assigned = assigned_by_index.get(i, 0.0)
        assigned_fv = assigned * (((1 - tax_rate) * (1 + yld) ** g["years"]) + tax_rate)
        remaining_target = max(0.0, infl - assigned_fv)
        sip = (
            0.0
            if remaining_target <= 1e-8
            else required_sip(remaining_target, g["years"], yld, tax_rate)
        )
        lumpsum = (
            0.0
            if remaining_target <= 1e-8
            else required_lumpsum(remaining_target, g["years"], yld, tax_rate)
        )
        out_goals.append({**g, "monthlySip": sip, "lumpsum": lumpsum, "assigned": assigned})

    return {
        "goals": out_goals,
        "totalMonthlySip": sum(g["monthlySip"] for g in out_goals),
        "totalLumpsum": sum(g["lumpsum"] for g in out_goals),
        "totalAssigned": sum(g["assigned"] for g in out_goals),
        "unassignedCorpus": remaining,
    }


def close(a: float, b: float, rel: float = 1e-8) -> bool:
    return abs(a - b) <= max(1e-4, abs(b) * rel)


def main() -> int:
    if not XLSX.exists():
        print(f"Missing workbook: {XLSX}", file=sys.stderr)
        return 1

    wb = load_workbook(XLSX, data_only=False)
    assert "Goal Calculator" in wb.sheetnames
    assert "GCwithCorpus" in wb.sheetnames

    ws = wb["Goal Calculator"]
    assert ws["A2"].value and "Goal Calculator" in str(ws["A2"].value)
    assert float(ws["D6"].value) == 0.07
    assert float(ws["D7"].value) == 0.15
    assert ws["G6"].value == 5
    assert ws["J7"].value == 12
    assert ws["B15"].value == "Education"
    assert ws["C15"].value == 5_000_000
    assert ws["D15"].value == 10
    assert ws["B21"].value == "Car"
    assert ws["C21"].value == 4_800_000
    print("OK Goal Calculator title + ST/LT + sample goals")

    sample = model(
        short_term_years=5,
        short_term_return=0.07,
        long_term_return=0.15,
        inflation=0.0,
        tax_rate=0.0,
        current_corpus=0.0,
        goals=[
            {"name": "Education", "amount": 5_000_000, "years": 10},
            {"name": "House1", "amount": 8_000_000, "years": 8},
            {"name": "House2", "amount": 130_000_000, "years": 12},
            {"name": "Car", "amount": 4_800_000, "years": 5},
            {"name": "Marriage", "amount": 50_000_000, "years": 25},
        ],
    )
    expect = {
        "eduSip": 19010.09236480497,
        "carSip": 67040.59538205592,
        "totalMonthlySip": 495205.7565509509,
        "totalLumpsum": 33090282.767640818,
    }
    assert close(sample["goals"][0]["monthlySip"], expect["eduSip"]), sample["goals"][0]
    assert close(sample["goals"][3]["monthlySip"], expect["carSip"]), sample["goals"][3]
    assert close(sample["totalMonthlySip"], expect["totalMonthlySip"])
    assert close(sample["totalLumpsum"], expect["totalLumpsum"])
    print(f"OK engine-model totalMonthlySip={sample['totalMonthlySip']}")

    gc = wb["GCwithCorpus"]
    assert gc["M6"].value == 100_000_000
    assert abs(float(gc["D6"].value) - 0.07) < 1e-12
    assert abs(float(gc["D7"].value) - 0.12) < 1e-12
    assert abs(float(gc["D8"].value) - 0.03) < 1e-12
    assert abs(float(gc["D9"].value) - 0.125) < 1e-12
    print("OK GCwithCorpus corpus + rate inputs")

    with_corpus = model(
        short_term_years=5,
        short_term_return=0.07,
        long_term_return=0.12,
        inflation=0.03,
        tax_rate=0.125,
        current_corpus=100_000_000,
        goals=[
            {"name": "Education", "amount": 4_000_000, "years": 6},
            {"name": "Masters", "amount": 440_000_000, "years": 11},
            {"name": "House", "amount": 2_000_000_000, "years": 10},
            {"name": "Marriage", "amount": 4_800_000, "years": 5},
            {"name": "Retirement", "amount": 50_000_000, "years": 25},
        ],
    )
    assert close(with_corpus["totalAssigned"], 100_000_000)
    assert with_corpus["unassignedCorpus"] == 0
    assert with_corpus["goals"][0]["assigned"] > 0  # Education funded first-ish by years
    assert with_corpus["goals"][3]["assigned"] > 0  # Marriage (5y) soonest
    assert with_corpus["goals"][0]["monthlySip"] == 0
    assert with_corpus["goals"][3]["monthlySip"] == 0
    print(f"OK engine-model corpus assign totalAssigned={with_corpus['totalAssigned']}")

    print("All audits passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
