#!/usr/bin/env python3
"""
Audit Unprotected Nivra Periodic Investment v1.xlsm and assert formula parity.

  .venv/bin/python calculator-tests/audit_periodic_investment.py
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Install openpyxl first: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

XLSX = Path(__file__).resolve().parent / "Nivra Periodic Investment v1.xlsm"
SHEET = "Periodic"


def monthly_rate(annual: float) -> float:
    return (1 + annual) ** (1 / 12) - 1


def fv_begin(rate: float, nper: float, pv: float) -> float:
    """Excel FV(rate, nper, 0, -pv, 1) with pmt=0."""
    return pv * (1 + rate) ** nper


def model(amount: float, times_per_year: int, years: int, annual_return: float, tax_rate: float):
    if times_per_year <= 0 or 12 % times_per_year != 0:
        raise ValueError("times_per_year must divide 12")
    interval = 12 / times_per_year
    total_months = years * 12
    r = monthly_rate(annual_return)
    schedule = []
    maturity = 0.0
    invested = 0.0
    payments = 0
    for month in range(total_months):
        if month % interval == 0:
            payments += 1
            invested += amount
            contribution_fv = fv_begin(r, total_months - month, amount)
            maturity += contribution_fv
            schedule.append(
                {
                    "month": month,
                    "contribution": amount,
                    "contributionFv": contribution_fv,
                }
            )
    gain = maturity - invested
    tax = max(0.0, gain) * tax_rate
    return {
        "maturity": maturity,
        "totalInvested": invested,
        "gain": gain,
        "tax": tax,
        "netAfterTax": maturity - tax,
        "payments": payments,
        "schedule": schedule,
    }


def close(a: float, b: float, rel: float = 1e-9) -> bool:
    return abs(a - b) <= max(1e-6, abs(b) * rel)


def main() -> int:
    if not XLSX.exists():
        print(f"Missing workbook: {XLSX}", file=sys.stderr)
        return 1

    wb = load_workbook(XLSX, data_only=False, keep_vba=True)
    assert SHEET in wb.sheetnames, wb.sheetnames
    ws = wb[SHEET]

    print("=== Periodic Investment v1 — field audit ===\n")
    editable = {
        "C5": "Name",
        "C6": "Age",
        "C8": "Periodic Investment Amount",
        "C9": "Freq. of Periodic Investment per year",
        "C10": "Term (in Years)",
        "C11": "Expected Rate of Return",
        "C12": "Capital Gains Tax Rate",
    }
    print("Editable inputs:")
    for addr, label in editable.items():
        print(f"  {addr}: {label} = {ws[addr].value!r} locked={ws[addr].protection.locked}")

    print("\nComputed (formulas):")
    for addr in ["C14", "C15", "C17", "C19", "C20", "O8", "O10", "O11", "P11"]:
        print(f"  {addr}: {ws[addr].value!r}")

    print("\nSchedule row pattern (contribution months only on web):")
    print("  M13.. = month index 0..")
    print("  N13 = IF(M<$O$8, IF(MOD(M,12/freq)=0, amount, 0), 0)")
    print("  O13 = FV(SRate, totalMonths-M, 0, -N, 1) when M<=totalMonths")

    cached = load_workbook(XLSX, data_only=True, keep_vba=True)[SHEET]
    sample = model(100_000, 2, 1, 0.12, 0.12)
    checks = [
        ("C14", sample["maturity"]),
        ("C15", sample["gain"]),
        ("C17", sample["tax"]),
        ("C19", sample["totalInvested"]),
        ("C20", sample["netAfterTax"]),
        ("O10", sample["totalInvested"]),
        ("O11", sample["maturity"]),
        ("P11", sample["tax"]),
    ]
    print("\n=== Sample parity (₹1L × 2 / yr, 1y, 12%, 12% tax) ===")
    failed = 0
    for addr, expected in checks:
        actual = cached[addr].value
        ok = actual is not None and close(float(actual), float(expected))
        status = "OK" if ok else "FAIL"
        if not ok:
            failed += 1
        print(f"  {status} {addr}: excel={actual!r} model={expected}")

    # Contribution months in cached sheet (N>0)
    contrib_months = []
    for r in range(13, 13 + 12):
        n = cached.cell(r, 14).value  # N
        m = cached.cell(r, 13).value  # M
        o = cached.cell(r, 15).value  # O
        if n and float(n) > 0:
            contrib_months.append((int(m), float(n), float(o)))
    print("\nExcel contribution rows (year 1):")
    for month, contrib, cfv in contrib_months:
        print(f"  Month {month}: contrib={contrib} fv={cfv}")
    expected_months = [row["month"] for row in sample["schedule"]]
    actual_months = [m for m, _, _ in contrib_months]
    if actual_months != expected_months:
        print(f"  FAIL schedule months excel={actual_months} model={expected_months}")
        failed += 1
    else:
        print("  OK contribution months match model")

    print("\n=== Frequency schedule length checks (model) ===")
    for freq, label in [(12, "Monthly"), (4, "Quarterly"), (2, "Half-Yearly"), (1, "Yearly")]:
        r = model(100_000, freq, 1, 0.12, 0.12)
        months = [row["month"] for row in r["schedule"]]
        print(f"  {label}: {r['payments']} payments → months {months}")
        assert r["payments"] == freq
        assert r["payments"] == len(r["schedule"])

    if failed:
        print(f"\n{failed} check(s) failed", file=sys.stderr)
        return 1
    print("\nAll checks passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
